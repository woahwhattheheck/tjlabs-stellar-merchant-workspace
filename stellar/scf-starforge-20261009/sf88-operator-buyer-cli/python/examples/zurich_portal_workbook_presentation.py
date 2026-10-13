#!/usr/bin/env python3
"""Build and receive a readable successor of the original Zurich customer workbook.

The complete original customer archive is the source. Saved queries, typed
results, definitions and deliveries keep their identities and histories.
"""
from __future__ import annotations
import argparse
from hashlib import sha256
from itertools import zip_longest
import json
import os
from pathlib import Path
import time
import uuid

from stellar_forge_buyer.analysis_workspace import AnalysisWorkspace
from stellar_forge_buyer.analysis_workbooks import WorkbookDelivery, _publish_received_directory
from stellar_forge_buyer.storage import dump, parse, private_directory, read
from stellar_forge_buyer.workspace_archive import inspect_archive
from stellar_forge_buyer.workspace_archive_receive import restore_archive

OPERATION = "LIIP-CKAN-WORKSHARE-814B-20261013"
PHASE = "original-workbook-presentation-v1"
SOURCE_ARCHIVE_SHA256 = "8023daaa5f2d8f9fb40a55bf222dadcff20da9d3ed4aa944fa7677b259e4243f"
ORIGINAL_WORKBOOK = "0320865c3c7452a5b88e7bb5a4a02155"
WORKSPACE_ID = "6118448bd4ef45acb85835c06790e37b"


def require(value, message):
    if not value:
        raise RuntimeError(message)


def identifier(label):
    return uuid.uuid5(uuid.NAMESPACE_URL, OPERATION + ":" + PHASE + ":" + label).hex


def file_record(path):
    import hashlib
    with Path(path).open("rb") as stream:
        return {"path": str(Path(path).absolute()), "bytes": os.fstat(stream.fileno()).st_size,
                "sha256": hashlib.file_digest(stream, "sha256").hexdigest()}


def keep(path, source):
    """Publish a fully received file without leaving a partial final pathname."""
    import shutil
    import stat
    import tempfile
    path, source = Path(path), Path(source)
    expected = file_record(source)

    def existing():
        info = path.lstat()
        require(stat.S_ISREG(info.st_mode) and info.st_uid == os.getuid() and not info.st_mode & 0o077,
                "Private original file required: " + path.name)
        actual = file_record(path)
        require((actual["bytes"], actual["sha256"]) == (expected["bytes"], expected["sha256"]),
                "Retained visual delivery changed: " + path.name)
        return actual

    if os.path.lexists(path):
        return existing()
    fd, name = tempfile.mkstemp(prefix=".receiving-", dir=path.parent)
    staging = Path(name)
    try:
        with source.open("rb") as incoming, os.fdopen(fd, "wb") as output:
            os.fchmod(output.fileno(), 0o600)
            shutil.copyfileobj(incoming, output, 1024 * 1024)
            output.flush()
            os.fsync(output.fileno())
        staged = file_record(staging)
        require((staged["bytes"], staged["sha256"]) == (expected["bytes"], expected["sha256"]),
                "Original output changed while receiving")
        try:
            # The existing native rename primitive also atomically admits a
            # regular file, avoiding a hard-link alias if the process dies.
            _publish_received_directory(staging, path)
        except FileExistsError:
            return existing()
    finally:
        staging.unlink(missing_ok=True)
    from stellar_forge_buyer.storage import sync_directory
    sync_directory(path.parent)
    return existing()

def settle(workspace, kind, job):
    deadline = time.monotonic() + 900
    while job["status"] == "RUNNING" and time.monotonic() < deadline:
        time.sleep(0.05)
        job = workspace.job(job["id"])
    require(job["status"] == "COMPLETE", kind + ": " + str(job.get("error") or job["status"]))
    print(json.dumps({"operation": OPERATION, "phase": PHASE, "kind": kind,
                      "id": job["id"], "status": job["status"]}), flush=True)
    return job


def compare_original_cells(original, successor, logical_names):
    from openpyxl import load_workbook
    before = load_workbook(original, read_only=True, data_only=False)
    after = load_workbook(successor, read_only=True, data_only=False)
    sentinel, counts = object(), {}
    try:
        require(after.sheetnames == ["Readme", "Daily requests", "Resource comparison", "Source observations", "Sources"],
                "Readable opening and data-sheet order changed")
        require(after.active.title == "Readme", "Workbook must open on its original delivery context")
        for name in logical_names:
            count = 0
            for left, right in zip_longest(before[name].iter_rows(values_only=True),
                                          after[name].iter_rows(values_only=True), fillvalue=sentinel):
                require(left is not sentinel and right is not sentinel and left == right,
                        "Original typed workbook cells changed in " + name)
                count += 1
            counts[name] = count - 1
        require(counts["Source observations"] == 8210, "Both complete original daily captures are required")
    finally:
        before.close()
        after.close()
    return counts


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--archive", required=True, help="Original complete 98958 Zurich customer workspace ZIP")
    parser.add_argument("--out", required=True, help="One private successor operation directory")
    parser.add_argument("--retry-workbook", action="store_true",
                        help="Save a new child attempt after a failed original build, preserving its history")
    args = parser.parse_args()
    os.umask(0o077)
    out = private_directory(Path(args.out).absolute(), create=True)
    source_archive = Path(args.archive).absolute()
    source = file_record(source_archive)
    require(source["sha256"] == SOURCE_ARCHIVE_SHA256, "Select the complete original 98958 customer archive")
    database = out / "workspace"
    record_path = out / "workbook-presentation-programme.json"
    binding = {"operation": OPERATION, "phase": PHASE, "sourceSha256": SOURCE_ARCHIVE_SHA256,
               "out": str(out), "database": str(database)}
    record = parse(read(record_path, 16 * 1024 * 1024)) if record_path.exists() else {
        "schema": "stellar-forge.zurich-workbook-presentation.v1",
        "binding": binding, "state": "RUNNING"}
    require(record["binding"] == binding, "Continue the same original presentation operation")
    restore_archive(source_archive, database, operation=identifier("source-receive"),
                    expected_sha256=SOURCE_ARCHIVE_SHA256)
    workspace = AnalysisWorkspace(database)
    try:
        require(workspace.identity == WORKSPACE_ID, "Original customer workspace identity changed")
        original_job = workspace.job(ORIGINAL_WORKBOOK)
        require(original_job["status"] == "COMPLETE", "The completed original workbook is required")
        original = WorkbookDelivery(database / "analysis-workbooks" / ORIGINAL_WORKBOOK)
        original_manifest = original.inspect()
        saved = workspace.workbooks.definition("ZurichPortalUseWorkbook", revision=1)
        require(saved["definition_sha256"] == original_manifest["definition_sha256"],
                "Original workbook definition and delivery differ")
        report_definition_sha = workspace.reports.definition("ZurichPortalUse")["sha256"]
        by_name = {sheet["name"]: sheet for sheet in saved["definition"]["sheets"]}
        require(set(by_name) == {"Source observations", "Resource comparison", "Daily requests"},
                "Original three workbook selections changed")
        next_definition = workspace.workbooks.save("ZurichPortalUseWorkbook", {
            "id": identifier("definition"), "expected_revision": 1,
            "title": saved["definition"]["title"],
            "notes": saved["definition"]["notes"] + "\n\nWorkbook reading order: Readme, daily requests, full resource comparison, and all original source observations. Filters and frozen headings remain available on every data sheet. Long original values wrap within their cells.",
            "sheets": [by_name[name] for name in ("Daily requests", "Resource comparison", "Source observations")]})
        request = {"id": record.get("workbookRequestId", identifier("build")),
                   "name": next_definition["name"], "revision": next_definition["revision"],
                   "definition_sha256": next_definition["definition_sha256"]}
        record["workbookRequestId"] = request["id"]
        dump(record_path, record)
        job = workspace.start("workbook", request)
        if job["status"] in {"FAILED", "CANCELLED", "INTERRUPTED"} and args.retry_workbook:
            attempts = record.setdefault("buildAttempts", [])
            if not any(old["id"] == job["id"] for old in attempts):
                attempts.append({"id": job["id"], "status": job["status"], "error": job.get("error")})
            request["id"] = identifier("retry:" + job["id"])
            record["workbookRequestId"] = request["id"]
            dump(record_path, record)
            job = workspace.start("workbook", request)
        built = settle(workspace, "workbook", job)
        delivery = WorkbookDelivery(database / "analysis-workbooks" / built["id"])
        manifest = delivery.inspect()
        require(manifest["excel"].get("presentation", {}).get("version") == 1,
                "The maintained readable workbook generator must be selected")
        original_sheets = {sheet["name"]: sheet for sheet in original_manifest["sheets"]}
        for sheet in manifest["sheets"]:
            require(sheet == original_sheets[sheet["name"]], "Original typed source projection changed")
        row_counts = compare_original_cells(original.directory / "workbook.xlsx",
                                           delivery.directory / "workbook.xlsx", by_name)
        handoff = delivery.directory / "handoff.zip"
        retained_handoff = keep(out / "Zurich_Workbook_Complete_Handoff.zip", handoff)
        received_path = out / "workbook-received"
        if received_path.exists():
            received = WorkbookDelivery(received_path)
        else:
            received = WorkbookDelivery.receive(handoff, received_path,
                                                 expected_sha256=retained_handoff["sha256"])
        require(received.inspect() == manifest, "Received workbook changed")
        editable = keep(out / "Zurich_Portal_Use_Workbook.xlsx", received_path / "workbook.xlsx")
        require(workspace.workbooks.definition("ZurichPortalUseWorkbook", revision=1)["definition_sha256"]
                == saved["definition_sha256"], "Original immutable workbook revision changed")
        require(workspace.reports.definition("ZurichPortalUse")["sha256"] == report_definition_sha,
                "Original published report changed")
        require(workspace.archives.review()["ready"], "The complete successor workspace must be exportable")
        export = settle(workspace, "workspace-export", workspace.start("workspace-export", {
            "id": identifier("workspace-export"), "include_previous_backups": False}))
        exported = keep(out / "Zurich_Customer_Workspace.zip",
                        Path(export["result"]["directory"]) / "workspace.zip")
        inspection = inspect_archive(Path(exported["path"]), expected_sha256=exported["sha256"])
        record.update({"workspaceId": workspace.identity, "originalWorkbookId": ORIGINAL_WORKBOOK,
                       "originalDefinitionSha256": saved["definition_sha256"],
                       "workbookId": built["id"], "definitionSha256": next_definition["definition_sha256"],
                       "revision": next_definition["revision"], "rows": row_counts,
                       "originalReportSha256": report_definition_sha,
                       "workbook": editable, "workbookHandoff": retained_handoff,
                       "workspaceArchive": {**exported, "payloadFiles": len(inspection["manifest"]["files"])}})
        dump(record_path, record)
    except Exception as error:
        record.update(state="INCOMPLETE", error={"type": type(error).__name__, "message": str(error)})
        dump(record_path, record)
        raise
    finally:
        workspace.close()
    cold = out / "customer-received"
    receiving = restore_archive(Path(record["workspaceArchive"]["path"]), cold,
                                operation=identifier("customer-receive"),
                                expected_sha256=record["workspaceArchive"]["sha256"])
    reopened = AnalysisWorkspace(cold)
    try:
        require(reopened.identity == WORKSPACE_ID, "Cold-received workspace identity changed")
        require(reopened.workbooks.definition("ZurichPortalUseWorkbook", revision=1)["definition_sha256"]
                == record["originalDefinitionSha256"], "Cold-received original revision changed")
        require(reopened.workbooks.definition("ZurichPortalUseWorkbook", revision=record["revision"])["definition_sha256"]
                == record["definitionSha256"], "Cold-received new revision changed")
        require(reopened.reports.definition("ZurichPortalUse")["sha256"] == record["originalReportSha256"],
                "Cold-received original report changed")
        for job_id in (ORIGINAL_WORKBOOK, record["workbookId"]):
            require(reopened.job(job_id)["status"] == "COMPLETE", "Complete workbook history is required")
            WorkbookDelivery(cold / "analysis-workbooks" / job_id).inspect()
        record["coldReopen"] = {"operationId": receiving["operation_id"], "workspaceId": reopened.identity,
                                "originalAndNewWorkbooks": True, "originalReport": True}
    finally:
        reopened.close()
    record.update(state="COMPLETE", renderedInspection="awaiting-producing-peer-view")
    record.pop("error", None)
    dump(record_path, record)
    print(json.dumps(record, ensure_ascii=False), flush=True)


if __name__ == "__main__":
    main()
